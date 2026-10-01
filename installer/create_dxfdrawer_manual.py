from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\MyDisk\Projects\DXF-explorer")
OUT = ROOT / "docs" / "DXFDrawer_사용자_매뉴얼.docx"
SCREENSHOT = Path(r"C:\Users\boyho\AppData\Local\Temp\codex-clipboard-5f6b8fee-f927-4e25-bebb-dea6b7ce2534.png")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Cm(widths[i])
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2:
                set_cell_shading(cell, "F3F7FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(str(value))
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def step(doc, number, title, detail):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    r = p.add_run(f"{number}. {title}  ")
    r.bold = True
    p.add_run(detail)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Malgun Gothic"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.12
for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 10.5)):
    style = styles[name]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("DXFDrawer 사용자 매뉴얼")
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("DXF 도면 작성 편집 저장 및 절곡 단면 확인")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(70, 70, 70)

doc.add_paragraph()
if SCREENSHOT.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SCREENSHOT), width=Cm(17.5))
    cap = doc.add_paragraph("DXFDrawer 기본 작업 화면")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8.5)

intro = doc.add_paragraph()
intro.add_run("이 매뉴얼의 목적  ").bold = True
intro.add_run("DXFDrawer를 처음 사용하는 사용자가 화면 구성을 이해하고, 기존 DXF 파일을 열어 수정하거나 새로운 DXF 도면을 작성하여 저장할 수 있도록 안내합니다. 프로그램은 LINE, ARC, CIRCLE 형상과 L, V, V1 레이어를 중심으로 판금 도면을 편집합니다.")

add_table(doc, ["문서 항목", "내용"], [
    ("대상 프로그램", "DXFDrawer.exe"),
    ("지원 작업", "DXF 열기, 신규 작성, 형상 편집, Validation 확인, DXF 저장"),
    ("주요 도형", "LINE, ARC, CIRCLE"),
    ("주요 레이어", "-L- 절단, -V- 정방향 절곡, -V1- 반대방향 절곡"),
], [4.2, 13.2])

doc.add_page_break()
heading(doc, "1 화면 구성", 1)
doc.add_paragraph("DXFDrawer의 기본 화면은 상단 도구 모음, 왼쪽 설정 및 객체 목록, 중앙 Drawing Canvas, 오른쪽과 아래쪽의 절곡 단면 시뮬레이션 및 Property Editor, 하단 상태 표시줄로 구성됩니다.")
add_table(doc, ["영역", "명칭", "기능"], [
    ("A", "상단 도구 모음", "파일 작업, 도형 도구 선택, 실행 취소·다시 실행, 삭제, 화면 맞춤을 실행합니다."),
    ("B", "Layer Visibility", "L, V, V1 레이어를 화면에서 개별 표시하거나 숨깁니다. 숨김은 표시만 바꾸며 객체를 삭제하지 않습니다."),
    ("C", "Grid 및 Snap 설정", "격자 표시, 격자점 스냅, 끝점 스냅과 격자 간격을 설정합니다."),
    ("D", "Contour 및 Objects", "전체 Entity 수와 닫힌 윤곽 수를 표시하고, 객체 목록에서 도형을 선택·검색합니다."),
    ("E", "Drawing Canvas", "도면을 작성하고 도형·핸들·단면선을 마우스로 편집하는 주 작업 영역입니다."),
    ("F", "Y Section Bending Simulation", "세로 Y Section 선이 통과하는 단면의 절곡 상태를 시뮬레이션합니다."),
    ("G", "X Section Bending Simulation", "가로 X Section 선이 통과하는 단면의 절곡 상태를 시뮬레이션합니다."),
    ("H", "Property Editor", "선택한 객체의 레이어, 좌표, 길이, 각도, 중심, 반경 등을 숫자로 편집합니다."),
    ("I", "상태 표시줄", "마우스 좌표, 현재 도구·레이어, 선택 객체, X/Y Section 위치와 저장 결과를 표시합니다."),
], [1.0, 4.5, 12.0])

heading(doc, "1.1 화면 분할 조정", 2)
doc.add_paragraph("Drawing Canvas와 오른쪽 영역 사이의 세로 구분선, 위·아래 영역 사이의 가로 구분선을 드래그하면 각 패널의 크기를 조절할 수 있습니다.")

heading(doc, "1.2 Drawing Canvas 마우스 조작", 2)
bullet(doc, "보기: 마우스 휠로 확대·축소하고, 오른쪽 버튼 드래그로 도면 화면을 이동합니다.")
bullet(doc, "선택과 찾기: 객체를 클릭하면 주황색 핸들이 표시됩니다. Objects 항목을 더블 클릭하면 해당 객체가 Canvas 중앙에 옵니다.")
bullet(doc, "취소와 삭제: Esc는 진행 중인 작성을 취소해 원래 상태로 복원하고, Delete는 선택 객체를 삭제합니다.")

heading(doc, "2 상단 도구 모음", 1)
add_table(doc, ["명칭", "기능과 사용 방법"], [
    ("새로 만들기", "현재 도면을 비우고 새 문서를 시작합니다. 새 문서의 저장 파일명은 아직 지정되지 않습니다."),
    ("열기", "DXF 파일 선택 창에서 기존 .dxf 파일을 불러옵니다."),
    ("저장", "Validation을 통과한 현재 도면을 DXF 파일로 저장합니다. 저장할 때마다 파일 선택 창이 열립니다."),
    ("선택", "객체 선택, 이동, 끝점·중심·반경·원호 각도 핸들 편집에 사용합니다."),
    ("선", "-L- 절단 레이어에 직선을 작성합니다."),
    ("원호", "-L- 절단 레이어에 반원을 작성합니다. 첫 점은 중심, 두 번째 점은 시작점과 반경을 정합니다."),
    ("원", "-L- 절단 레이어에 원을 작성합니다. 첫 점은 중심, 두 번째 점은 반경을 정합니다."),
    ("V-Line", "-V- 정방향 절곡 레이어에 수평 또는 수직 직선을 작성합니다."),
    ("V1-Line", "-V1- 반대방향 절곡 레이어에 수평 또는 수직 직선을 작성합니다."),
    ("실행 취소", "직전 변경 전의 도면으로 되돌립니다."),
    ("다시 실행", "실행 취소한 변경을 다시 적용합니다."),
    ("삭제", "선택한 객체를 삭제합니다."),
    ("맞춤", "전체 도면이 Drawing Canvas 안에 보이도록 확대 비율과 위치를 맞춥니다."),
], [4.0, 13.5])

heading(doc, "2.1 도형 작성의 공통 방식", 2)
doc.add_paragraph("작성 도구를 선택한 후 Drawing Canvas에서 시작 위치를 누르고 원하는 위치까지 드래그한 뒤 버튼을 놓습니다. 도형이 생성되면 프로그램은 자동으로 선택 도구로 돌아갑니다.")
bullet(doc, "Grid Snap은 포인터를 설정한 Grid 간격에 맞춥니다.")
bullet(doc, "Endpoint Snap은 가까운 기존 끝점에 우선 연결하며, V-Line과 V1-Line은 수평 또는 수직 방향으로 자동 제한됩니다.")

heading(doc, "3 레이어 객체 목록 및 속성 편집", 1)
heading(doc, "3.1 레이어와 표시 색상", 2)
add_table(doc, ["레이어", "용도", "화면 색상", "허용 도형"], [
    ("-L-", "절단 형상", "검정", "LINE, ARC, CIRCLE"),
    ("-V-", "정방향 절곡선", "빨강", "LINE만 허용"),
    ("-V1-", "반대방향 절곡선", "파랑", "LINE만 허용"),
], [2.4, 5.0, 3.0, 7.0])
doc.add_paragraph("Layer Visibility의 체크를 끄면 해당 레이어만 화면에서 숨겨집니다. 객체와 저장 데이터는 유지됩니다.")

heading(doc, "3.2 Objects 목록", 2)
doc.add_paragraph("Objects에는 객체 번호, 형식, 레이어가 표시됩니다. Validation 오류가 있는 객체에는 경고 기호가 붙습니다. 목록을 클릭하면 Canvas의 같은 객체가 선택되고, 더블 클릭하면 해당 객체가 화면 중앙에 위치합니다.")

heading(doc, "3.3 Property Editor", 2)
add_table(doc, ["객체", "편집 가능 항목", "읽기 전용 정보"], [
    ("LINE", "Layer, Start X/Y, End X/Y, Length, Angle °", "Entity, ΔX, ΔY, Gradient, Direction"),
    ("CIRCLE", "Layer, Center X/Y, Radius", "Entity"),
    ("ARC", "Layer, Center X/Y, Radius, Start Angle, End Angle", "Entity, Sweep"),
], [3.0, 9.2, 5.2])
doc.add_paragraph("값을 수정한 뒤 Enter를 누르면 즉시 도면에 반영됩니다. 각도는 도면 좌표계에서 도(degree) 단위이며, 0°는 오른쪽, 90°는 위쪽입니다. 원호의 Start Angle과 End Angle은 0° 이상 360° 미만으로 정규화됩니다.")

heading(doc, "3.4 선택 핸들 편집", 2)
bullet(doc, "LINE: 두 끝점 핸들을 드래그해 시작점 또는 끝점을 변경합니다. 선 본체를 드래그하면 전체 선을 이동합니다.")
bullet(doc, "CIRCLE: 중심 핸들을 드래그하면 중심을 이동하고, 원 둘레를 드래그하면 반경을 변경합니다.")
bullet(doc, "ARC: 중심 핸들은 중심 이동, 시작·끝 핸들은 각도 변경, 원호 본체는 반경 변경에 사용합니다.")

doc.add_page_break()
heading(doc, "4 기존 DXF 파일 열기와 수정", 1)
step(doc, 1, "DXFDrawer 실행", "시작 메뉴 또는 설치 폴더의 DXFDrawer.exe를 실행합니다.")
step(doc, 2, "파일 열기", "상단의 열기를 누르고 수정할 .dxf 파일을 선택합니다.")
step(doc, 3, "화면 확인", "도면이 자동으로 화면에 맞춰지고, Objects의 Entity 수와 Closed contour 수가 갱신되는지 확인합니다.")
step(doc, 4, "객체 찾기", "Canvas에서 객체를 클릭하거나 Objects 목록에서 항목을 선택합니다. 찾기 어려우면 목록 항목을 더블 클릭합니다.")
step(doc, 5, "형상 수정", "선택 핸들을 드래그하거나 Property Editor에 좌표·반경·각도 값을 입력하고 Enter를 누릅니다.")
step(doc, 6, "추가 또는 삭제", "필요하면 선, 원호, 원, V-Line, V1-Line 도구로 객체를 추가합니다. 불필요한 객체는 선택 후 삭제 또는 Delete 키로 제거합니다.")
step(doc, 7, "검증", "Objects의 경고 표시와 화면 메시지를 확인합니다. 오류 객체를 수정한 뒤 저장합니다.")
step(doc, 8, "저장", "저장을 누르고 원본을 유지하려면 다른 파일명을 지정합니다. 현재 파일명이 기본값으로 제시됩니다.")

heading(doc, "4.1 안전한 수정 권장 순서", 2)
bullet(doc, "원본 보존이 필요하면 먼저 다른 이름으로 저장할 파일명을 계획합니다.")
bullet(doc, "Grid와 Endpoint Snap을 작업 정밀도에 맞게 켜거나 끕니다.")
bullet(doc, "숫자로 정확히 수정해야 하는 값은 Property Editor를 사용합니다.")
bullet(doc, "오조작 시 실행 취소를 사용하고, 결과가 맞으면 맞춤으로 전체 형상을 다시 확인합니다.")

heading(doc, "4.2 지원 범위", 2)
doc.add_paragraph("현재 편집 모델은 LINE, ARC, CIRCLE 객체를 사용합니다. 파일에 지원하지 않는 Entity가 포함되어 있으면 열기 오류 메시지가 표시될 수 있습니다. 저장 결과는 DXF AC1009 형식의 ENTITIES 섹션에 형상과 레이어 정보를 기록합니다.")

doc.add_page_break()
heading(doc, "5 새로운 DXF 파일 작성과 저장", 1)
step(doc, 1, "새 문서 시작", "새로 만들기를 누릅니다. Drawing Canvas와 Objects 목록이 비워집니다.")
step(doc, 2, "격자 설정", "Grid 표시와 Grid Snap을 선택하고 Grid (mm)에 필요한 간격을 입력한 뒤 다른 곳을 클릭합니다. 기본값은 10 mm입니다.")
step(doc, 3, "절단 외곽 작성", "선, 원호, 원 도구를 사용해 -L- 레이어의 절단 형상을 작성합니다. 연결점은 Endpoint Snap으로 정확히 맞춥니다.")
step(doc, 4, "절곡선 작성", "필요한 방향에 따라 V-Line 또는 V1-Line을 사용해 수평·수직 절곡선을 작성합니다.")
step(doc, 5, "정밀 치수 입력", "각 객체를 선택해 Property Editor에서 좌표, 길이, 각도, 중심, 반경을 입력하고 Enter를 누릅니다.")
step(doc, 6, "윤곽 확인", "왼쪽의 Closed contour 수를 확인합니다. 절단 외곽의 끝점이 서로 연결되어 닫힌 윤곽을 이루는지 검토합니다.")
step(doc, 7, "단면 검토", "X Section과 Y Section 선을 필요한 위치로 드래그하고 단면 시뮬레이션에서 절곡 위치를 확인합니다.")
step(doc, 8, "DXF 저장", "저장을 누르고 파일명과 저장 위치를 지정합니다. 새 문서의 기본 파일명은 drawing.dxf입니다.")

heading(doc, "5.1 직선 작성", 2)
doc.add_paragraph("선을 누른 뒤 시작점에서 끝점까지 드래그합니다. 작성 후 끝점 핸들 또는 Property Editor의 Start X/Y, End X/Y, Length, Angle °로 치수를 조정할 수 있습니다.")

heading(doc, "5.2 원과 원호 작성", 2)
doc.add_paragraph("원과 원호는 첫 클릭·드래그 시작점이 중심입니다. 드래그를 놓는 지점까지의 거리가 반경이 됩니다. 원호는 해당 방향을 시작점으로 하는 180° 원호로 생성되며, 이후 시작·끝 핸들 또는 Start Angle과 End Angle로 필요한 호 길이를 설정합니다.")

heading(doc, "5.3 절곡선 작성", 2)
doc.add_paragraph("V-Line은 빨간색 정방향 절곡, V1-Line은 파란색 반대방향 절곡입니다. 두 도구 모두 수평 또는 수직선으로 자동 제한됩니다. 절단선과 같은 구간에 겹치지 않도록 작성합니다.")

doc.add_page_break()
heading(doc, "6 X Y Section 절곡 시뮬레이션", 1)
doc.add_paragraph("Drawing Canvas의 녹색 점선은 X Section, 자홍색 점선은 Y Section입니다. 점선 가까이를 왼쪽 버튼으로 드래그하면 단면 위치가 바뀌고 해당 시뮬레이션이 즉시 갱신됩니다.")
add_table(doc, ["구성", "기능"], [
    ("X Section Bending Simulation", "가로 단면선이 통과하는 절단 형상과 V/V1 절곡점을 아래 패널에 표시합니다."),
    ("Y Section Bending Simulation", "세로 단면선이 통과하는 절단 형상과 V/V1 절곡점을 오른쪽 패널에 표시합니다."),
    ("빨간 절곡점", "-V- 정방향 절곡 Entity입니다."),
    ("파란 절곡점", "-V1- 반대방향 절곡 Entity입니다."),
    ("절곡점 클릭", "선택한 절곡을 접거나 펼쳐 단면 형상 변화를 확인합니다."),
    ("Reset", "해당 단면의 모든 절곡 시뮬레이션 상태를 펼쳐진 초기 상태로 되돌립니다."),
], [6.0, 11.5])
doc.add_paragraph("단면선이 절곡선을 통과하지 않으면 No bending points on this section이 표시됩니다. 단면 형상을 만들 수 없으면 No section geometry가 표시될 수 있습니다.")

doc.add_page_break()
heading(doc, "7 Validation과 저장 제한", 1)
doc.add_paragraph("DXFDrawer는 잘못된 형상을 저장하지 않도록 작성·편집 중과 저장 전에 Validation을 수행합니다. 오류가 있으면 저장을 중단하고 첫 오류 객체를 선택합니다.")
add_table(doc, ["오류 유형", "의미와 해결 방법"], [
    ("유효하지 않은 좌표", "좌표가 숫자가 아니거나 유한한 값이 아닙니다. Property Editor에서 정상 숫자로 수정합니다."),
    ("길이가 0인 선", "시작점과 끝점이 같습니다. 한쪽 끝점을 이동하거나 Length를 0보다 크게 입력합니다."),
    ("잘못된 절곡 방향", "V/V1 선이 대각선입니다. 수평 또는 수직이 되도록 좌표를 수정합니다."),
    ("잘못된 원 또는 원호", "Radius가 0 이하이거나 중심·각도 값이 유효하지 않습니다. Radius와 좌표를 수정합니다."),
    ("절단선과 절곡선 중첩", "-L- 선과 -V-/-V1- 선이 동일 구간에 겹칩니다. 둘 중 하나를 이동하거나 길이를 조정합니다."),
], [5.2, 12.3])

doc.add_page_break()
heading(doc, "8 빠른 작업 예시", 1)
heading(doc, "8.1 기존 원호의 시작점 변경", 2)
step(doc, 1, "원호 선택", "Canvas에서 원호를 클릭하거나 Objects의 ARC 항목을 더블 클릭합니다.")
step(doc, 2, "시작점 이동", "주황색 시작점 핸들을 원호 중심을 기준으로 원하는 각도 위치까지 드래그합니다.")
step(doc, 3, "수치 확인", "Property Editor의 Center X/Y, Radius, Start Angle, End Angle과 Sweep를 확인합니다.")
step(doc, 4, "정확한 각도 입력", "필요하면 Start Angle에 예를 들어 90을 입력하고 Enter를 누릅니다. 중심과 반경은 유지되고 시작점만 중심 기준 90° 위치로 이동합니다.")

heading(doc, "8.2 직사각형 절단 외곽 작성", 2)
step(doc, 1, "Grid 준비", "Grid Snap과 Endpoint Snap을 켜고 필요한 Grid 간격을 입력합니다.")
step(doc, 2, "네 변 작성", "선 도구로 네 개의 -L- 직선을 순서대로 작성하며 각 끝점을 다음 선의 시작점에 맞춥니다.")
step(doc, 3, "닫힌 윤곽 확인", "Closed contour가 1인지 확인합니다. 0이면 Objects에서 선을 선택해 끝점 좌표를 맞춥니다.")
step(doc, 4, "절곡선 추가", "V-Line 또는 V1-Line으로 내부 절곡선을 추가하고 X/Y Section 시뮬레이션을 확인합니다.")
step(doc, 5, "저장", "저장을 눌러 파일명을 지정합니다.")

heading(doc, "9 문제 해결", 1)
add_table(doc, ["증상", "확인 및 조치"], [
    ("객체가 보이지 않음", "Layer Visibility가 꺼져 있는지 확인하고 맞춤을 누릅니다."),
    ("원하는 좌표로 이동하지 않음", "Grid Snap 또는 Endpoint Snap의 영향을 확인합니다. 정밀 값은 Property Editor에 입력합니다."),
    ("V/V1 선이 대각선으로 되지 않음", "정상 동작입니다. 절곡선은 수평 또는 수직으로 제한됩니다."),
    ("저장되지 않음", "Validation 오류 메시지와 경고 표시 객체를 확인하고 수정합니다."),
    ("단면에 절곡점이 없음", "X/Y Section 선이 V/V1 선을 실제로 통과하는 위치인지 확인합니다."),
    ("편집을 취소하고 싶음", "Esc를 눌러 진행 중인 드래그를 취소하거나 실행 취소를 사용합니다."),
], [5.2, 12.3])

heading(doc, "10 작업 전 확인 사항", 1)
bullet(doc, "중요 파일은 원본과 다른 이름으로 저장합니다.")
bullet(doc, "절단 외곽의 연결과 Closed contour 수를 확인합니다.")
bullet(doc, "V와 V1 레이어 방향이 설계 의도와 맞는지 색상과 시뮬레이션으로 확인합니다.")
bullet(doc, "Objects의 경고 표시가 없고 저장이 정상 완료되었는지 상태 표시줄에서 확인합니다.")

for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DXFDrawer 사용자 매뉴얼")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 100, 100)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "DXFDrawer 사용자 매뉴얼"
doc.core_properties.subject = "DXF 도면 작성 편집 저장 및 절곡 단면 확인"
doc.core_properties.author = "DXFExplorer"
doc.save(OUT)
print(OUT)
