from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "DXFExplorer_설치_매뉴얼.docx"
ICON = ROOT / "cnc_vgroove_icon.png"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, value=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    repeat_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(cell)
        shade(cell, "1F4E78")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_no, values in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(values):
            cell = cells[i]
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            if row_no % 2:
                shade(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(str(text))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, level=0):
    return doc.add_paragraph(text, style="List Bullet" if level == 0 else "List Bullet 2")


def step(doc, number, title, detail):
    p = doc.add_paragraph()
    r = p.add_run(f"{number}. {title}  ")
    r.bold = True
    p.add_run(detail)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.59)
section.page_height = Cm(27.94)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(1.9)
section.right_margin = Cm(1.9)

styles = doc.styles
styles["Normal"].font.name = "Malgun Gothic"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.13
for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 11)):
    style = styles[name]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("DXFExplorer 설치 매뉴얼")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("DXFExplorerSetup 설치 실행 구성 및 제거 안내")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(70, 70, 70)

if ICON.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ICON), width=Cm(5.0))

p = doc.add_paragraph()
p.add_run("이 매뉴얼의 목적  ").bold = True
p.add_run("DXFExplorerSetup.exe를 사용해 DXFExplorer 제품군을 Windows에 설치하고, 설치 후 생성되는 프로그램과 위치를 확인하며, 필요할 때 안전하게 제거할 수 있도록 안내합니다.")

add_table(doc, ["항목", "내용"], [
    ("설치 파일", "DXFExplorerSetup.exe"),
    ("제품 버전", "1.1.0"),
    ("지원 환경", "Windows x64 호환 환경"),
    ("기본 설치 경로", r"C:\Program Files\DXFExplorer"),
    ("권한", "설치와 제거 시 관리자 권한 필요"),
    ("설치 언어", "한국어 또는 English"),
], [4.2, 13.0])

doc.add_page_break()
heading(doc, "설치 결과 요약", 1)
doc.add_paragraph(r"설치가 끝나면 C:\Program Files\DXFExplorer 디렉터리에 다음 네 개의 사용자 프로그램이 생성됩니다. 각 프로그램은 같은 아이콘을 사용하며 시작 메뉴의 DXFExplorer 그룹에서도 실행할 수 있습니다.")
add_table(doc, ["실행 파일", "기능"], [
    ("DXFExplorer.exe", "DXF 파일을 찾아보고 제품군의 주요 기능으로 접근하는 파일 탐색 프로그램입니다."),
    ("DXFViewer.exe", "DXF 도면과 레이어 및 형상을 읽고 화면에 표시하는 뷰어입니다."),
    ("DXFSimulator.exe", "DXF 도면의 가공 및 절곡 관련 정보를 시뮬레이션하는 프로그램입니다."),
    ("DXFDrawer.exe", "DXF 도면을 새로 작성하거나 기존 도면을 열어 편집하고 저장하는 프로그램입니다."),
], [5.0, 12.2])

doc.add_page_break()
heading(doc, "1 설치 전 준비", 1)
bullet(doc, "Windows x64 호환 PC에서 설치합니다.")
bullet(doc, "설치 중인 프로그램 파일을 교체할 수 있도록 실행 중인 DXFExplorer, DXFViewer, DXFSimulator, DXFDrawer를 닫습니다.")
bullet(doc, "설치 파일의 전체 이름과 위치를 확인합니다. 기본 배포 파일명은 DXFExplorerSetup.exe입니다.")
bullet(doc, "Program Files에 설치하므로 관리자 계정 또는 관리자 승인 수단을 준비합니다.")
bullet(doc, "기존 도면 파일은 설치 과정에서 삭제되지 않지만, 중요한 DXF 파일은 별도로 백업하는 것이 좋습니다.")

heading(doc, "2 설치 방법", 1)
step(doc, 1, "설치 파일 실행", "DXFExplorerSetup.exe를 더블 클릭합니다.")
step(doc, 2, "사용자 계정 컨트롤 승인", "Windows가 이 앱이 장치를 변경하도록 허용할지 묻는 경우 예를 선택합니다.")
step(doc, 3, "설치 언어 선택", "한국어 또는 English를 선택합니다. 이후 설치 마법사의 버튼과 안내가 선택한 언어로 표시됩니다.")
step(doc, 4, "설치 시작", "환영 화면의 안내를 읽고 다음을 누릅니다.")
step(doc, 5, "설치 위치 확인", r"기본 경로 C:\Program Files\DXFExplorer를 확인합니다. 특별한 이유가 없다면 기본 경로를 사용합니다.")
step(doc, 6, "추가 작업 선택", "바탕 화면에 DXFExplorer 바로가기를 만들려면 해당 항목을 선택합니다. 이 항목은 기본적으로 선택되어 있지 않습니다.")
step(doc, 7, "설치 실행", "설치 준비 내용을 확인하고 설치를 누릅니다. 프로그램과 자체 포함형 .NET 지원 파일이 복사됩니다.")
step(doc, 8, "설치 완료", "마침 화면에서 DXFExplorer 실행을 선택하면 마침을 누른 직후 DXFExplorer.exe가 시작됩니다.")

heading(doc, "2.1 설치 중 표시될 수 있는 안내", 2)
doc.add_paragraph("설치 프로그램은 현재 실행 중인 관련 프로그램을 닫도록 안내할 수 있습니다. 작업 내용을 먼저 저장한 다음 프로그램을 종료하고 설치를 계속합니다. 설치는 재부팅을 기본으로 요구하지 않습니다.")

doc.add_page_break()
heading(doc, "3 설치되는 디렉터리와 파일", 1)
heading(doc, "3.1 기본 설치 디렉터리", 2)
p = doc.add_paragraph()
p.add_run(r"C:\Program Files\DXFExplorer").bold = True
p.add_run("에 제품군의 실행 파일, 프로그램 라이브러리, .NET 실행 지원 파일과 언어별 리소스가 함께 설치됩니다.")

heading(doc, "3.2 사용자가 실행하는 파일", 2)
add_table(doc, ["전체 경로", "용도"], [
    (r"C:\Program Files\DXFExplorer\DXFExplorer.exe", "파일 탐색 및 제품군 진입"),
    (r"C:\Program Files\DXFExplorer\DXFViewer.exe", "DXF 도면 보기"),
    (r"C:\Program Files\DXFExplorer\DXFSimulator.exe", "DXF 시뮬레이션"),
    (r"C:\Program Files\DXFExplorer\DXFDrawer.exe", "DXF 신규 작성 및 편집"),
], [11.2, 6.0])

heading(doc, "3.3 지원 파일", 2)
add_table(doc, ["파일 또는 폴더", "설명"], [
    ("*.dll", "각 프로그램과 WPF 및 .NET 실행에 필요한 라이브러리입니다."),
    ("*.deps.json", "프로그램별 종속성 정보를 기록합니다."),
    ("*.runtimeconfig.json", "프로그램이 사용할 .NET 런타임 설정입니다."),
    ("coreclr.dll, hostfxr.dll, hostpolicy.dll 등", "별도 .NET 설치 없이 실행하기 위한 자체 포함형 런타임 파일입니다."),
    ("언어 코드 폴더", "ko, en 및 기타 언어별 WPF 리소스가 들어 있습니다."),
    ("createdump.exe", "오류 진단을 위한 .NET 런타임 지원 도구이며 사용자가 직접 실행할 파일이 아닙니다."),
    ("unins000.exe 및 unins000.dat", "설치 후 생성되는 제거 프로그램과 제거 정보입니다. 번호는 설치 환경에 따라 달라질 수 있습니다."),
], [6.8, 10.4])
doc.add_paragraph("지원 파일을 개별 삭제하거나 다른 위치로 옮기면 프로그램이 실행되지 않을 수 있습니다. 설치 폴더의 파일은 설치 프로그램 또는 Windows의 제거 기능으로 관리합니다.")

heading(doc, "3.4 바로가기", 2)
bullet(doc, "시작 메뉴의 DXFExplorer 그룹: DXFExplorer, DXFViewer, DXFSimulator, DXFDrawer와 DXFExplorer 제거 바로가기가 생성됩니다.")
bullet(doc, "바탕 화면: 설치 중 추가 작업을 선택한 경우 DXFExplorer 바로가기 하나가 생성됩니다.")
bullet(doc, "바로가기를 삭제해도 실제 프로그램 파일은 삭제되지 않습니다.")

doc.add_page_break()
heading(doc, "4 설치 후 확인", 1)
step(doc, 1, "설치 폴더 확인", r"파일 탐색기에서 C:\Program Files\DXFExplorer를 엽니다.")
step(doc, 2, "네 실행 파일 확인", "DXFExplorer.exe, DXFViewer.exe, DXFSimulator.exe, DXFDrawer.exe가 모두 있는지 확인합니다.")
step(doc, 3, "시작 메뉴 확인", "시작 메뉴에서 DXFExplorer 그룹을 찾아 네 프로그램의 바로가기가 보이는지 확인합니다.")
step(doc, 4, "프로그램 실행", "각 프로그램을 한 번씩 실행해 기본 화면이 정상적으로 표시되는지 확인합니다.")
step(doc, 5, "도면 파일 확인", "업무에 사용하는 DXF 파일을 복사본으로 열어 표시 또는 편집 기능을 점검합니다.")

heading(doc, "5 프로그램 실행 방법", 1)
add_table(doc, ["방법", "절차"], [
    ("시작 메뉴", "Windows 시작에서 DXFExplorer 그룹을 열고 원하는 프로그램을 선택합니다."),
    ("바탕 화면", "설치할 때 바로가기를 선택했다면 DXFExplorer 아이콘을 더블 클릭합니다."),
    ("설치 폴더", r"C:\Program Files\DXFExplorer에서 원하는 .exe 파일을 직접 실행합니다."),
], [4.6, 12.6])

heading(doc, "6 업데이트 또는 재설치", 1)
doc.add_paragraph("같은 제품의 새 설치 파일을 실행하면 기존 설치 위치를 재사용하고 실행 파일과 지원 파일을 갱신할 수 있습니다. 업데이트 전 네 프로그램을 모두 종료하고 중요한 도면 작업을 저장합니다.")
bullet(doc, "설치 프로그램의 제품 식별자는 동일하게 유지되므로 기존 설치를 같은 제품으로 인식합니다.")
bullet(doc, "이전에 선택한 설치 경로와 시작 메뉴 그룹이 제안될 수 있습니다.")
bullet(doc, "업데이트 후 문제가 있으면 Windows 설정에서 기존 제품을 제거한 뒤 새 설치 파일로 다시 설치합니다.")

doc.add_page_break()
heading(doc, "7 제거 방법", 1)
heading(doc, "7.1 Windows 설정에서 제거", 2)
step(doc, 1, "프로그램 종료", "실행 중인 네 프로그램을 모두 닫습니다.")
step(doc, 2, "설치된 앱 열기", "Windows 설정에서 앱, 설치된 앱을 엽니다.")
step(doc, 3, "제품 찾기", "DXFExplorer를 검색합니다.")
step(doc, 4, "제거 실행", "제거를 선택하고 관리자 권한 요청을 승인합니다.")
step(doc, 5, "완료 확인", r"제거가 끝난 뒤 C:\Program Files\DXFExplorer 폴더와 시작 메뉴 바로가기가 정리되었는지 확인합니다.")

heading(doc, "7.2 시작 메뉴에서 제거", 2)
doc.add_paragraph("시작 메뉴의 DXFExplorer 그룹에서 DXFExplorer 제거를 선택한 뒤 안내에 따라 진행할 수도 있습니다.")

heading(doc, "7.3 사용자 DXF 파일", 2)
doc.add_paragraph("설치 제거는 설치 프로그램이 배치한 프로그램 파일을 대상으로 합니다. 사용자가 별도 폴더에 작성하거나 저장한 DXF 파일은 제거 대상이 아닙니다. 단, 중요한 업무 파일은 제거 전에 위치를 확인하고 백업합니다.")

heading(doc, "8 문제 해결", 1)
add_table(doc, ["증상", "확인 및 해결"], [
    ("관리자 권한 요청이 나타남", "정상 동작입니다. Program Files에 설치하기 위해 예를 선택합니다."),
    ("파일 사용 중 메시지가 나타남", "실행 중인 네 프로그램을 종료한 다음 다시 시도합니다."),
    ("바탕 화면 아이콘이 없음", "바탕 화면 바로가기는 선택 사항입니다. 시작 메뉴 또는 설치 폴더에서 실행합니다."),
    ("프로그램 실행 파일이 없음", "설치 폴더가 올바른지 확인하고 설치 프로그램을 다시 실행합니다."),
    ("DLL 관련 오류가 나타남", "설치 폴더의 지원 파일이 삭제되었을 수 있습니다. 재설치하여 파일을 복원합니다."),
    ("제거할 수 없음", "프로그램을 모두 종료하고 관리자 권한으로 Windows 설정의 제거를 다시 실행합니다."),
], [5.2, 12.0])

heading(doc, "9 설치 완료 확인표", 1)
bullet(doc, r"C:\Program Files\DXFExplorer 폴더가 생성되었습니다.")
bullet(doc, "DXFExplorer.exe, DXFViewer.exe, DXFSimulator.exe, DXFDrawer.exe가 모두 존재합니다.")
bullet(doc, "시작 메뉴의 DXFExplorer 그룹에서 네 프로그램을 실행할 수 있습니다.")
bullet(doc, "선택한 경우 바탕 화면에 DXFExplorer 바로가기가 있습니다.")
bullet(doc, "각 프로그램의 기본 화면이 오류 없이 열립니다.")

for paragraph in list(doc.paragraphs) + [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]:
    for run in paragraph.runs:
        if "\\" in run.text:
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")

for sec in doc.sections:
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DXFExplorer 설치 매뉴얼")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 100, 100)

doc.core_properties.title = "DXFExplorer 설치 매뉴얼"
doc.core_properties.subject = "DXFExplorerSetup 설치 실행 구성 및 제거 안내"
doc.core_properties.author = "DXFExplorer"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
